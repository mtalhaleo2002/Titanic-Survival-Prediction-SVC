
from pathlib import Path
import subprocess
import sys
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
SCREENSHOT_DIR = BASE_DIR / "screenshots"
OUTPUT_DIR = BASE_DIR / "output"

SCREENSHOT_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

WORD_FILE = OUTPUT_DIR / "Titanic_Assignment.docx"

TASK_TITLES = {
    1: "Dataset Loading and Initial View",
    2: "Data Exploration",
    3: "Data Cleaning",
    4: "Feature Selection and SVC Model",
    5: "Data Visualization",
}

TASK_DESCRIPTIONS = {
    1: (
        "Loaded the Titanic dataset using Pandas and "
        "displayed the dataset dimensions and first seven rows."
    ),
    2: (
        "Explored numerical and categorical variables, "
        "calculated descriptive statistics, and analyzed "
        "gender and passenger class frequency distributions."
    ),
    3: "Identified and handled missing values in the dataset.",
    4: (
        "Prepared features, trained an SVC classification "
        "model, and evaluated its performance."
    ),
    5: (
        "Visualized survival patterns based on gender, "
        "passenger class, and fare."
    ),
}


def create_images(content, filename_prefix, title):
    """Render genuine file contents/output into PNG pages."""

    font_path = (
        "/usr/share/fonts/truetype/dejavu/"
        "DejaVuSansMono.ttf"
    )

    font = ImageFont.truetype(font_path, 17)
    title_font = ImageFont.truetype(font_path, 22)

    lines = content.splitlines() or [""]

    # Separate long content into readable images.
    max_lines = 34
    image_paths = []

    for start in range(0, len(lines), max_lines):
        chunk = lines[start:start + max_lines]
        part = start // max_lines + 1

        width = 1500
        line_height = 30
        height = 110 + len(chunk) * line_height

        img = Image.new("RGB", (width, height), "#1e1e1e")
        draw = ImageDraw.Draw(img)

        draw.rectangle(
            (0, 0, width, 65),
            fill="#252526"
        )

        draw.text(
            (30, 18),
            f"{title} - Part {part}",
            font=title_font,
            fill="#ffffff"
        )

        y = 85

        for line in chunk:
            draw.text(
                (30, y),
                line[:125],
                font=font,
                fill="#d4d4d4"
            )
            y += line_height

        suffix = "" if part == 1 else f"_part{part}"

        image_path = SCREENSHOT_DIR / (
            f"{filename_prefix}{suffix}.png"
        )

        img.save(image_path)
        image_paths.append(image_path)

        print(f"Generated: {image_path.name}")

    return image_paths


def execute_task(task_file):
    """Run task using the current virtual environment."""

    result = subprocess.run(
        [sys.executable, str(task_file)],
        cwd=BASE_DIR,
        capture_output=True,
        text=True,
        timeout=120
    )

    output = result.stdout

    if result.stderr:
        output += "\nSTDERR:\n" + result.stderr

    if result.returncode != 0:
        raise RuntimeError(
            f"{task_file.name} failed:\n{output}"
        )

    return output


def add_images(doc, image_paths):
    for image_path in image_paths:
        doc.add_picture(
            str(image_path),
            width=Inches(6.7)
        )


def create_report():

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    doc.add_heading(
        "Artificial Intelligence Using Python",
        0
    )

    doc.add_paragraph(
        "DigiSkills 3.0 - Batch 04\n"
        "Hands-on Exercise No. 02\n"
        "Titanic Dataset Analysis"
    )

    task_files = sorted(
        BASE_DIR.glob("task[1-5].py"),
        key=lambda path: int(
            re.search(r"task(\d+)", path.stem).group(1)
        )
    )

    if not task_files:
        raise FileNotFoundError(
            "No task1.py - task5.py files found!"
        )

    for task_file in task_files:

        task_number = int(
            re.search(r"task(\d+)", task_file.stem).group(1)
        )

        print(f"\nProcessing Task {task_number}...")

        code_content = task_file.read_text(
            encoding="utf-8"
        )

        output_content = execute_task(task_file)

        code_images = create_images(
            code_content,
            f"task{task_number}_code",
            f"VS Code - {task_file.name}"
        )

        output_images = create_images(
            output_content,
            f"task{task_number}_output",
            f"Terminal - Task {task_number} Output"
        )

        doc.add_heading(
            f"Task {task_number}: "
            f"{TASK_TITLES[task_number]}",
            level=1
        )

        doc.add_paragraph(
            TASK_DESCRIPTIONS[task_number]
        )

        doc.add_heading("Python Code", level=2)
        add_images(doc, code_images)

        doc.add_heading("Execution Output", level=2)
        add_images(doc, output_images)

        # Include graphs generated by each task.
        graph_files = sorted(
            OUTPUT_DIR.glob(f"task{task_number}_graph*.png")
        )

        if graph_files:
            doc.add_heading("Visualizations", level=2)

            for graph_file in graph_files:
                doc.add_picture(
                    str(graph_file),
                    width=Inches(6.4)
                )

        doc.add_paragraph(
            f"Task {task_number} executed successfully."
        )

        print(f"Task {task_number} added to Word!")

    doc.save(WORD_FILE)

    print("\n" + "=" * 55)
    print("ALL AVAILABLE TASKS PROCESSED SUCCESSFULLY")
    print("=" * 55)
    print(f"Word document: {WORD_FILE}")


if __name__ == "__main__":
    create_report()
